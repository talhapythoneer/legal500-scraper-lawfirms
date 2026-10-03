#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Legal 500 Law Firm Scraper
==========================

Scrapes law firm profiles from https://www.legal500.com organised by the
site's own geographic hierarchy (Region > Sub-Region > Country > Law firms)
for the regions: Europe, Middle East, Africa, Americas and Asia Pacific.

For every firm it collects:
    - Practice areas (with Legal 500 ranking tier)   <- client-requested
    - About / firm description                        <- client-requested
    - Office address(es)                               <- client-requested
    - Listing (profile) URL
    - Firm name, Country, Sub-Region, Region
    - Website
    - Phone number(s) and email(s) per office
    - Lawyers (name + position)
    - Which extra editorial sections the firm has (News & developments,
      Client testimonials, Diversity, Teams, Interviews, Comparative guides)

Output: a single .xlsx workbook with one sheet per region. Inside each
sheet, rows are sorted by Sub-Region > Country > Firm name, which mirrors
the geographic hierarchy the client asked for
(e.g. Americas / Latin America / Argentina > law firms).

Column order: the three fields the client explicitly asked for (Practice
Areas, About, Address(es)) plus the Listing URL come FIRST, exactly as
requested. Every other field scraped comes afterwards.

WHY curl_cffi AND NOT SELENIUM/PLAYWRIGHT
------------------------------------------
legal500.com is a Next.js site sitting behind Cloudflare. It renders fully
on the server (no client-side data fetching needed for these pages) and its
anti-bot check is TLS/JA3 fingerprint-based, not a JS challenge. A plain
`requests`/`scrapy` call already returns full 200-OK HTML with no
JS-challenge page, and `curl_cffi` (which impersonates a real Chrome TLS/
HTTP2 fingerprint) is used here purely for extra robustness/longevity
against Cloudflare, while staying lightweight and fast - no browser,
no driver binaries, far higher throughput than Selenium/Playwright would
give for ~4,000+ firms x ~3-5 requests each.

USAGE
-----
    python legal500_scraper.py                       # full run, all 5 regions
    python legal500_scraper.py --countries portugal,spain
    python legal500_scraper.py --regions americas,europe
    python legal500_scraper.py --limit 20             # quick smoke test
    python legal500_scraper.py --workers 16
    python legal500_scraper.py --build-only           # just rebuild the
                                                       # .xlsx from the cache,
                                                       # no network calls

The scraper is resumable: every firm record is appended to a local JSONL
cache (`legal500_cache.jsonl` by default) as soon as it's scraped. Re-running
the script skips firms already present in the cache, so an interrupted run
(or a crash) can simply be restarted with the same command.

Requirements: curl_cffi, lxml, openpyxl  (all already used in this project).
"""

from __future__ import annotations

import argparse
import json
import logging
import os
import re
import sys
import threading
import time
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from typing import Optional

from curl_cffi import requests as cfr
from lxml import html
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #

BASE_URL = "https://www.legal500.com"

# The site groups "/rankings" into these top-level region ids. We only want
# the regions the client asked for (this excludes "United Kingdom" and
# "Offshore Britain", which the site also tracks as separate top-level
# regions).
TARGET_REGION_IDS = ["europe", "middle-east", "africa", "americas", "asia-pacific"]

REQUEST_TIMEOUT = 40
MAX_RETRIES = 3
RETRY_BACKOFF = 1.6  # seconds, multiplied by attempt number

# Final column order in the workbook. The client's requested fields plus the
# listing URL come first (per instructions); everything else follows.
COLUMNS = [
    "Practice Areas",
    "About",
    "Address(es)",
    "Listing URL",
    "Firm Name",
    "Country",
    "Sub-Region",
    "Region",
    "Website",
    "Phone Number(s)",
    "Email(s)",
    "Lawyers",
    "Other Sections Available",
]

# Columns that benefit from wrapped text / wider width in the workbook.
WRAP_COLUMNS = {"Practice Areas", "About", "Address(es)", "Lawyers"}

LOG = logging.getLogger("legal500_scraper")


# --------------------------------------------------------------------------- #
# HTTP helpers
# --------------------------------------------------------------------------- #

_thread_local = threading.local()


def get_session() -> "cfr.Session":
    """One curl_cffi Session per worker thread (keep-alive, not shared)."""
    sess = getattr(_thread_local, "session", None)
    if sess is None:
        sess = cfr.Session(impersonate="chrome")
        _thread_local.session = sess
    return sess


def fetch(url: str) -> Optional[str]:
    """GET a URL with retries. Returns decoded HTML text, or None on failure."""
    session = get_session()
    last_err = None
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            resp = session.get(url, timeout=REQUEST_TIMEOUT)
            if resp.status_code == 200:
                # The site serves UTF-8 but curl_cffi occasionally guesses a
                # different codec from headers; force UTF-8 decoding of the
                # raw bytes to avoid mangled accented characters.
                return resp.content.decode("utf-8", errors="replace")
            if resp.status_code in (429, 500, 502, 503, 504):
                last_err = f"HTTP {resp.status_code}"
            else:
                LOG.warning("Non-retryable status %s for %s", resp.status_code, url)
                return None
        except Exception as exc:  # noqa: BLE001 - network errors are varied
            last_err = str(exc)
        time.sleep(RETRY_BACKOFF * attempt)
    LOG.warning("Failed to fetch %s after %d attempts: %s", url, MAX_RETRIES, last_err)
    return None


def abs_url(href: str) -> str:
    if href.startswith("http"):
        return href
    return BASE_URL + href


# --------------------------------------------------------------------------- #
# Cloudflare email de-obfuscation
# --------------------------------------------------------------------------- #

def decode_cfemail(cfemail_hex: str) -> str:
    """Decode a Cloudflare `data-cfemail` hex string into a plain email."""
    try:
        key = int(cfemail_hex[:2], 16)
        return "".join(
            chr(int(cfemail_hex[i : i + 2], 16) ^ key)
            for i in range(2, len(cfemail_hex), 2)
        )
    except Exception:  # noqa: BLE001
        return ""


def clean_text(s: Optional[str]) -> str:
    if not s:
        return ""
    return re.sub(r"[ \t]+", " ", re.sub(r"\r\n|\r", "\n", s)).strip()


# --------------------------------------------------------------------------- #
# Data model
# --------------------------------------------------------------------------- #

@dataclass
class CountryRef:
    region: str
    sub_region: str
    country: str
    slug: str


@dataclass
class FirmRef:
    name: str
    href: str  # relative, e.g. /firms/10036-abreu-advogados/c-portugal
    region: str
    sub_region: str
    country: str


# --------------------------------------------------------------------------- #
# Step 1: discover the Region > Sub-Region > Country hierarchy
# --------------------------------------------------------------------------- #

def discover_country_hierarchy(region_filter: Optional[set] = None) -> list[CountryRef]:
    text = fetch(f"{BASE_URL}/rankings")
    if not text:
        raise RuntimeError("Could not load /rankings - cannot build hierarchy")
    doc = html.fromstring(text)

    region_label_by_id = {}
    countries: list[CountryRef] = []
    seen_slugs: set[str] = set()  # first-seen sub-region wins (handles overlap
    # categories like "Global Bar" / "Greater China" in Asia Pacific)

    for section in doc.xpath('//section[starts-with(@id,"/r/")]'):
        region_id = section.get("id")[3:]  # strip "/r/"
        if region_id not in TARGET_REGION_IDS:
            continue
        if region_filter and region_id not in region_filter:
            continue
        h2 = section.xpath("./h2")
        region_label = h2[0].text_content().strip() if h2 else region_id
        region_label_by_id[region_id] = region_label

        current_sub_region = ""
        for el in section.iter("h3", "a"):
            if el.tag == "h3":
                current_sub_region = el.text_content().strip()
            else:
                href = el.get("href") or ""
                if not href.startswith("/c/"):
                    continue
                slug = href[3:].strip("/")
                if not slug or slug in seen_slugs:
                    continue
                seen_slugs.add(slug)
                countries.append(
                    CountryRef(
                        region=region_label,
                        sub_region=current_sub_region,
                        country=el.text_content().strip(),
                        slug=slug,
                    )
                )
    return countries


# --------------------------------------------------------------------------- #
# Step 2: list every firm in a country's directory
# --------------------------------------------------------------------------- #

def discover_firms_in_country(country: CountryRef) -> list[FirmRef]:
    text = fetch(f"{BASE_URL}/c/{country.slug}/directory")
    if not text:
        return []
    doc = html.fromstring(text)
    h2 = doc.xpath('//h2[normalize-space()="Directory"]')
    if not h2:
        # Country has no listed firms (e.g. very small jurisdictions) - not
        # an error, just nothing to scrape there.
        return []
    section = h2[0].getparent()
    firms = []
    for a in section.xpath('.//article/a[contains(@href,"/firms/")]'):
        h3 = a.xpath(".//h3")
        name = h3[0].text_content().strip() if h3 else ""
        href = a.get("href")
        if not name or not href:
            continue
        firms.append(
            FirmRef(
                name=name,
                href=href,
                region=country.region,
                sub_region=country.sub_region,
                country=country.country,
            )
        )
    return firms


# --------------------------------------------------------------------------- #
# Step 3: scrape a single firm's profile
# --------------------------------------------------------------------------- #

EXTRA_SECTION_LABELS = {
    "Comparative guides",
    "News and developments",
    "Client testimonials",
    "Diversity",
    "Teams",
    "Interview with…",
}


def parse_nav_tabs(doc) -> dict[str, str]:
    """Returns {tab label: absolute href} from the firm sub-navigation."""
    tabs = {}
    for a in doc.xpath('//ul[@aria-label="Firm Sub Navigation"]//a'):
        label = a.text_content().strip()
        href = a.get("href")
        if label and href:
            tabs[label] = abs_url(href)
    return tabs


def parse_header(doc) -> tuple[str, str]:
    """Returns (firm_name, website_url) from the shared page header."""
    h1 = doc.xpath("//h1")
    name = h1[0].text_content().strip() if h1 else ""
    website = ""
    if h1:
        # the website link sits in the same header block as the h1, as an
        # external (_blank) link
        container = h1[0]
        for _ in range(3):
            container = container.getparent()
            if container is None:
                break
            links = container.xpath('.//a[@target="_blank"]/@href')
            if links:
                website = links[0]
                break
    return name, website


def parse_practice_areas(doc) -> list[tuple[str, str]]:
    """Returns [(practice area name, tier label)] from a Rankings-tab page.

    Most ranking cards show a numbered tier badge (1-6+). Some instead carry
    a non-numeric "Firms to Watch" badge image, and a few carry no badge at
    all (the firm is simply listed in that area without a tier). We capture
    all three cases.
    """
    out = []
    for a in doc.xpath('//a[contains(@href,"/rankings/ranking/")]'):
        h3 = a.xpath(".//h3")
        if not h3:
            continue
        pa_name = h3[0].text_content().strip()
        tier_el = a.xpath('.//span[contains(@class,"typography-interface")]')
        if tier_el:
            tier = f"Tier {tier_el[-1].text_content().strip()}"
        else:
            badge_alt = a.xpath(".//img/@alt")
            tier = badge_alt[0].strip() if badge_alt else ""
        if pa_name:
            out.append((pa_name, tier))
    return out


def parse_about(doc) -> str:
    """Extracts the About tab's content.

    Most firms have their About text as a series of <p> paragraphs inside
    the data-swiftype-name="body" div. Some firms instead (or additionally)
    have that body empty and carry their real content in sibling blocks
    such as "Languages", "Staffing Figures", "Memberships", etc. - each a
    <h3> heading followed by a <ul><li> list. We walk the whole About
    section container (not just the body div) so neither case is missed.
    """
    body_divs = doc.xpath('//div[@data-swiftype-name="body"]')
    if not body_divs:
        return ""
    container = body_divs[0].xpath("ancestor::section[1]")
    container = container[0] if container else body_divs[0].getparent()
    parts = []
    for el in container.iter("p", "h3", "li"):
        txt = clean_text(el.text_content())
        if not txt or txt.startswith("Content supplied by"):
            continue
        if el.tag == "h3":
            parts.append(f"\n{txt}:")
        elif el.tag == "li":
            parts.append(f"- {txt}")
        else:
            parts.append(txt)
    return "\n".join(parts).strip()


def parse_contact(doc) -> list[dict]:
    """Returns a list of office dicts: {label, address, phone, email}."""
    main = doc.xpath("//main")
    if not main:
        return []
    offices = []
    for art in main[0].xpath(".//article"):
        label = art.xpath('.//div[contains(@class,"typography-eyebrow")]/text()')
        addr_lines = art.xpath(
            './/div[contains(@class,"typography-body-prop")][1]//p/text()'
        )
        tel = art.xpath('.//a[starts-with(@href,"tel:")]/text()')
        cfemail = art.xpath('.//span[@class="__cf_email__"]/@data-cfemail')
        if not (label or addr_lines or tel or cfemail):
            continue
        offices.append(
            {
                "label": label[0].strip() if label else "",
                "address": ", ".join(a.strip() for a in addr_lines if a.strip()),
                "phone": tel[0].strip() if tel else "",
                "email": decode_cfemail(cfemail[0]) if cfemail else "",
            }
        )
    return offices


def parse_lawyers(doc) -> list[tuple[str, str]]:
    main = doc.xpath("//main")
    if not main:
        return []
    out = []
    for art in main[0].xpath(".//article"):
        h3 = art.xpath(".//h3")
        if not h3:
            continue
        name = clean_text(h3[0].text_content())
        pos_el = art.xpath('.//div[contains(@class,"typography-interface-s")]')
        pos = clean_text(pos_el[0].text_content()) if pos_el else ""
        if name:
            out.append((name, pos))
    return out


def format_practice_areas(items: list[tuple[str, str]]) -> str:
    return "; ".join(f"{name} ({tier})" if tier else name for name, tier in items)


def format_offices(offices: list[dict]) -> str:
    parts = []
    for o in offices:
        bits = [b for b in [o["label"], o["address"]] if b]
        parts.append(" - ".join(bits))
    return " | ".join(parts)


def format_phones(offices: list[dict]) -> str:
    seen, out = set(), []
    for o in offices:
        if o["phone"] and o["phone"] not in seen:
            seen.add(o["phone"])
            out.append(o["phone"])
    return "; ".join(out)


def format_emails(offices: list[dict]) -> str:
    seen, out = set(), []
    for o in offices:
        if o["email"] and o["email"] not in seen:
            seen.add(o["email"])
            out.append(o["email"])
    return "; ".join(out)


def format_lawyers(items: list[tuple[str, str]]) -> str:
    return "; ".join(f"{name} ({pos})" if pos else name for name, pos in items)


def scrape_firm(firm: FirmRef) -> dict:
    """Scrapes one firm's profile across its available tabs."""
    listing_url = abs_url(firm.href)

    # The Contact tab exists on virtually every firm profile and its page
    # carries the shared sub-navigation, so one request gives us both the
    # office/contact data AND the list of which other tabs exist.
    contact_html = fetch(listing_url + "/contact")
    offices: list[dict] = []
    website = ""
    firm_name = firm.name
    tabs: dict[str, str] = {}

    if contact_html:
        doc = html.fromstring(contact_html)
        offices = parse_contact(doc)
        name_from_page, website = parse_header(doc)
        if name_from_page:
            firm_name = name_from_page
        tabs = parse_nav_tabs(doc)

    practice_areas: list[tuple[str, str]] = []
    rankings_href = tabs.get("Rankings")
    if rankings_href:
        rk_html = fetch(rankings_href)
        if rk_html:
            practice_areas = parse_practice_areas(html.fromstring(rk_html))

    about_text = ""
    about_href = tabs.get("About")
    if about_href:
        about_html = fetch(about_href)
        if about_html:
            about_text = parse_about(html.fromstring(about_html))

    lawyers: list[tuple[str, str]] = []
    lawyers_href = tabs.get("Lawyers")
    if lawyers_href:
        lw_html = fetch(lawyers_href)
        if lw_html:
            lawyers = parse_lawyers(html.fromstring(lw_html))

    other_sections = sorted(set(tabs.keys()) & EXTRA_SECTION_LABELS)

    return {
        "Practice Areas": format_practice_areas(practice_areas),
        "About": about_text,
        "Address(es)": format_offices(offices),
        "Listing URL": listing_url,
        "Firm Name": firm_name,
        "Country": firm.country,
        "Sub-Region": firm.sub_region,
        "Region": firm.region,
        "Website": website,
        "Phone Number(s)": format_phones(offices),
        "Email(s)": format_emails(offices),
        "Lawyers": format_lawyers(lawyers),
        "Other Sections Available": "; ".join(other_sections),
    }


# --------------------------------------------------------------------------- #
# Cache (JSONL) - makes the crawl resumable
# --------------------------------------------------------------------------- #

class Cache:
    def __init__(self, path: str):
        self.path = path
        self._lock = threading.Lock()
        self.rows: dict[str, dict] = {}
        if os.path.exists(path):
            with open(path, "r", encoding="utf-8") as fh:
                for line in fh:
                    line = line.strip()
                    if not line:
                        continue
                    try:
                        row = json.loads(line)
                        self.rows[row["Listing URL"]] = row
                    except Exception:  # noqa: BLE001
                        continue
        self._fh = open(path, "a", encoding="utf-8")

    def has(self, listing_url: str) -> bool:
        return listing_url in self.rows

    def add(self, row: dict) -> None:
        with self._lock:
            self.rows[row["Listing URL"]] = row
            self._fh.write(json.dumps(row, ensure_ascii=False) + "\n")
            self._fh.flush()

    def close(self):
        self._fh.close()


# --------------------------------------------------------------------------- #
# Excel output
# --------------------------------------------------------------------------- #

def build_workbook(rows: list[dict], output_path: str) -> None:
    by_region: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        by_region[row.get("Region", "Unknown")].append(row)

    wb = Workbook()
    wb.remove(wb.active)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill_font = Font(bold=True)
    wrap_align = Alignment(wrap_text=True, vertical="top")
    top_align = Alignment(vertical="top")

    # Preserve a stable, sensible sheet order.
    region_order = ["Europe", "Middle East", "Africa", "Americas", "Asia Pacific"]
    ordered_regions = [r for r in region_order if r in by_region] + [
        r for r in by_region if r not in region_order
    ]

    for region in ordered_regions:
        region_rows = sorted(
            by_region[region],
            key=lambda r: (r.get("Sub-Region", ""), r.get("Country", ""), r.get("Firm Name", "")),
        )
        sheet_name = region[:31]
        ws = wb.create_sheet(title=sheet_name)

        ws.append(COLUMNS)
        for col_idx in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_fill_font
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

        for row in region_rows:
            ws.append([row.get(col, "") for col in COLUMNS])

        # Column widths + wrap for long text fields.
        for col_idx, col_name in enumerate(COLUMNS, start=1):
            letter = get_column_letter(col_idx)
            if col_name in WRAP_COLUMNS:
                ws.column_dimensions[letter].width = 60
            elif col_name in ("Firm Name", "Country", "Website"):
                ws.column_dimensions[letter].width = 28
            else:
                ws.column_dimensions[letter].width = 22
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=col_idx).alignment = (
                    wrap_align if col_name in WRAP_COLUMNS else top_align
                )

    # Summary sheet
    summary = wb.create_sheet(title="Summary", index=0)
    summary.append(["Region", "Firms Scraped"])
    summary["A1"].font = header_fill_font
    summary["B1"].font = header_fill_font
    total = 0
    for region in ordered_regions:
        n = len(by_region[region])
        total += n
        summary.append([region, n])
    summary.append(["TOTAL", total])
    summary.column_dimensions["A"].width = 20
    summary.column_dimensions["B"].width = 16

    wb.save(output_path)
    LOG.info("Workbook written: %s (%d firms total)", output_path, total)


# --------------------------------------------------------------------------- #
# Orchestration
# --------------------------------------------------------------------------- #

def run(args: argparse.Namespace) -> None:
    cache = Cache(args.cache)

    if args.build_only:
        build_workbook(list(cache.rows.values()), args.output)
        cache.close()
        return

    region_filter = None
    if args.regions:
        wanted = {r.strip().lower().replace(" ", "-") for r in args.regions.split(",")}
        region_filter = wanted

    LOG.info("Discovering Region > Sub-Region > Country hierarchy ...")
    countries = discover_country_hierarchy(region_filter)

    if args.countries:
        wanted_slugs = {c.strip().lower() for c in args.countries.split(",")}
        countries = [c for c in countries if c.slug in wanted_slugs]

    LOG.info("%d countries to scan.", len(countries))

    LOG.info("Listing firms per country ...")
    all_firms: list[FirmRef] = []
    with ThreadPoolExecutor(max_workers=args.workers) as ex:
        futures = {ex.submit(discover_firms_in_country, c): c for c in countries}
        for i, fut in enumerate(as_completed(futures), start=1):
            c = futures[fut]
            try:
                firms = fut.result()
            except Exception as exc:  # noqa: BLE001
                LOG.warning("Failed listing firms for %s: %s", c.country, exc)
                firms = []
            all_firms.extend(firms)
            if i % 20 == 0 or i == len(countries):
                LOG.info("  listed %d/%d countries, %d firms so far", i, len(countries), len(all_firms))

    if args.limit:
        all_firms = all_firms[: args.limit]

    # Skip firms already in the cache (resumability).
    todo = [f for f in all_firms if not cache.has(abs_url(f.href))]
    LOG.info(
        "%d firms discovered total, %d already cached, %d to scrape now.",
        len(all_firms),
        len(all_firms) - len(todo),
        len(todo),
    )

    start = time.time()
    done = 0
    lock = threading.Lock()

    def worker(f: FirmRef):
        nonlocal done
        try:
            row = scrape_firm(f)
            cache.add(row)
        except Exception as exc:  # noqa: BLE001
            LOG.warning("Failed scraping %s (%s): %s", f.name, f.href, exc)
        with lock:
            done += 1
            if done % 25 == 0 or done == len(todo):
                elapsed = time.time() - start
                rate = done / elapsed if elapsed > 0 else 0
                LOG.info(
                    "  scraped %d/%d firms (%.1f firms/min)",
                    done,
                    len(todo),
                    rate * 60,
                )

    with ThreadPoolExecutor(max_workers=args.workers) as ex:
        futures = [ex.submit(worker, f) for f in todo]
        for fut in as_completed(futures):
            fut.result()  # re-raise anything unexpected

    cache.close()

    # Rebuild the cache's file handle closed above; reopen read-only view by
    # re-reading what we have in memory (already up to date).
    build_workbook(list(cache.rows.values()), args.output)


def parse_args(argv=None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Legal 500 law firm scraper")
    p.add_argument(
        "--regions",
        help="Comma-separated region filter, e.g. 'europe,americas'. "
        "Default: all of europe, middle-east, africa, americas, asia-pacific.",
    )
    p.add_argument(
        "--countries",
        help="Comma-separated country slugs to restrict to, e.g. 'portugal,spain'.",
    )
    p.add_argument("--limit", type=int, help="Cap the number of firms scraped (for testing).")
    p.add_argument("--workers", type=int, default=10, help="Concurrent worker threads (default 10).")
    p.add_argument(
        "--cache",
        default="legal500_cache.jsonl",
        help="Path to the resumable JSONL cache file (default: legal500_cache.jsonl).",
    )
    p.add_argument(
        "--output",
        default="legal500_law_firms.xlsx",
        help="Path to the output .xlsx workbook (default: legal500_law_firms.xlsx).",
    )
    p.add_argument(
        "--build-only",
        action="store_true",
        help="Skip scraping; just rebuild the .xlsx from the existing cache file.",
    )
    p.add_argument("-v", "--verbose", action="store_true", help="Verbose (DEBUG) logging.")
    return p.parse_args(argv)


def main(argv=None) -> None:
    args = parse_args(argv)
    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        datefmt="%H:%M:%S",
    )
    run(args)


if __name__ == "__main__":
    main()
